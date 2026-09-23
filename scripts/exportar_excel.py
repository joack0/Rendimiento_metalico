"""Exporta a Excel toda la información del dashboard de rendimiento metálico.

Uso:
    python3 scripts/exportar_excel.py [ruta_csv] [ruta_maestro_xlsx]

Genera dashboard/Rendimiento_metalico.xlsx. Los resúmenes y la carta de control son
fórmulas sobre la hoja "Ordenes": si se corrige una orden ahí, todo se recalcula.
"""
import os
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from preparar_datos import CSV_DEFAULT, MAESTRO_DEFAULT, RAIZ, cargar  # noqa: E402

SALIDA = Path(os.environ.get("XLSX_SALIDA", RAIZ / "dashboard" / "Rendimiento_metalico.xlsx"))
GRUPOS = ["B Hormigón", "Redondo liso", "Plana", "Ángulo", "Cuadrado", "Estrella", "Saferock", "Hexágono"]
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

ROJO = "C8102E"
GRAFITO = "2B2B2B"
F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_titulo = Font(name=F, size=14, bold=True, color=GRAFITO)
f_nota = Font(name=F, size=9, italic=True, color="6F6E69")
f_input = Font(name=F, size=10, bold=True, color="0000FF")
fill_head = PatternFill("solid", fgColor=ROJO)
fill_sub = PatternFill("solid", fgColor="F0EFEC")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_alerta = PatternFill("solid", fgColor="F8D7DA")
borde = Border(bottom=Side(style="thin", color="E1E0D9"))

PCT = "0.00%"
KG = "#,##0"
PP = '+0.00" pp";-0.00" pp";0.00" pp"'


def encabezado(ws, fila, textos, col0=1):
    for j, t in enumerate(textos):
        c = ws.cell(fila, col0 + j, t)
        c.font, c.fill = f_head, fill_head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 30


def anchos(ws, valores):
    for j, w in enumerate(valores, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def titulo(ws, texto, sub=None):
    ws["A1"] = texto
    ws["A1"].font = f_titulo
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f_nota


def fuente_arial(ws):
    for fila in ws.iter_rows():
        for c in fila:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic,
                              color=c.font.color)


def semana_iso(fecha):
    """Semana ISO 8601 (lunes a domingo): [año ISO, número de semana, etiqueta]."""
    y, w, _ = fecha.isocalendar()
    return [y, w, f"{y}-S{w:02d}"]


def etiqueta_semana(y, w):
    lunes = datetime.fromisocalendar(y, w, 1)
    return f"S{w:02d} {y} ({lunes.day} {MESES[lunes.month - 1].lower()})"


def main():
    ruta = Path(sys.argv[1]) if len(sys.argv) > 1 else CSV_DEFAULT
    ruta_maestro = Path(sys.argv[2]) if len(sys.argv) > 2 else MAESTRO_DEFAULT
    productos, ordenes, maestro, sin_maestro = cargar(ruta, ruta_maestro)
    ordenes.sort(key=lambda o: (o["ini"], o["o"]))

    wb = Workbook()

    # ------------------------------------------------------------------ Ordenes
    wo = wb.active
    wo.title = "Ordenes"
    cols = ["Orden", "Fecha inicio", "Fecha", "Año", "Mes", "Grupo", "Producto", "Material original (CSV)",
            "Formato (maestro)", "Código (maestro)", "Diámetro (mm)", "Ancho (mm)", "Espesor (mm)", "Lado (mm)",
            "Largo (m)", "Kg producidos", "Kg consumidos", "Rend. met.", "Meta (Rend_Met)", "Estado",
            "Pérdida (kg)", "Meta × Kg cons", "Año ISO", "Semana ISO", "Semana"]
    encabezado(wo, 1, cols)
    n = len(ordenes)
    for i, o in enumerate(ordenes, 2):
        p = productos[o["p"]]
        fila = [int(o["o"]) if o["o"].isdigit() else o["o"], o["ini"], o["ini"].date(),
                f"=YEAR(C{i})", f"=MONTH(C{i})", p["grupo"], o["p"], o["original"],
                p.get("formato"), p.get("codigo"),
                p["diametro"], p["ancho"], p["espesor"], p["lado"], p["largo"],
                o["kp"], o["kc"], f"=IFERROR(P{i}/Q{i},\"\")",
                (o["meta"] / 100) if o["meta"] else None,
                f"=IF(R{i}>1,\"A corregir\",\"Válida\")", f"=Q{i}-P{i}", f"=IF(S{i}=\"\",0,S{i}*Q{i})",
                *semana_iso(o["ini"])]
        for j, v in enumerate(fila, 1):
            c = wo.cell(i, j, v)
            c.font = f_base
        wo.cell(i, 2).number_format = "dd-mm-yyyy hh:mm"
        wo.cell(i, 3).number_format = "dd-mm-yyyy"
        for j in (16, 17, 21, 22):
            wo.cell(i, j).number_format = KG
        wo.cell(i, 18).number_format = PCT
        wo.cell(i, 19).number_format = "0.0%"
        if p.get("largo_asumido"):
            wo.cell(i, 15).comment = Comment("Largo asumido: no está en el maestro ni en el nombre.", "Dashboard")
    ult = n + 1
    tab = Table(displayName="Ordenes", ref=f"A1:{get_column_letter(len(cols))}{ult}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    wo.add_table(tab)
    wo.freeze_panes = "B2"
    anchos(wo, [11, 16, 11, 7, 6, 13, 44, 44, 14, 12, 10, 10, 10, 9, 9, 13, 13, 10, 10, 11, 11, 14, 9, 9, 11])
    wo.conditional_formatting.add(f"T2:T{ult}", CellIsRule(operator="equal", formula=['"A corregir"'], fill=fill_alerta))
    wo.conditional_formatting.add(f"R2:R{ult}", CellIsRule(operator="greaterThan", formula=["1"], fill=fill_alerta))

    # rangos absolutos usados en todas las fórmulas
    R = {k: f"Ordenes!${c}$2:${c}${ult}" for k, c in
         {"fecha": "C", "anio": "D", "mes": "E", "grupo": "F", "prod": "G", "kp": "P", "kc": "Q",
          "estado": "T", "metakc": "V", "anioiso": "W", "sem": "X"}.items()}

    # ------------------------------------------------ hojas de carta de control
    def hoja_carta(nombre, titulo_txt, periodos, tipo):
        ws = wb.create_sheet(nombre)
        titulo(ws, titulo_txt, "Carta I-MR sobre el rendimiento ponderado (Σ kg producidos / Σ kg consumidos) "
                                "de las órdenes válidas (≤ 100 %).")
        ws["A4"], ws["B4"] = "Grupo:", "Todos"
        ws["A5"], ws["B5"] = "Producto:", "Todos"
        for c in ("A4", "A5"):
            ws[c].font = f_bold
        for c in ("B4", "B5"):
            ws[c].font, ws[c].fill = f_input, fill_input
        ws["C4"] = "← elige un grupo de la lista (celda amarilla) o deja «Todos»"
        ws["C5"] = "← escribe el nombre exacto del producto (columna Producto de «Ordenes») o deja «Todos»"
        ws["C4"].font = ws["C5"].font = f_nota
        dv = DataValidation(type="list", formula1='"Todos,' + ",".join(GRUPOS) + '"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add("B4")
        crit_g = 'IF($B$4="Todos","*",$B$4)'
        crit_p = 'IF($B$5="Todos","*",$B$5)'

        # parámetros de la carta
        par = [("LC (rend. ponderado)", "=IFERROR(SUM(E{a}:E{b})/SUM(F{a}:F{b}),\"\")", PCT),
               ("MR̄ (promedio rango móvil)", "=IFERROR(AVERAGE(J{a}:J{b}),\"\")", PCT),
               ("σ = MR̄ / 1,128", "=IFERROR(G8/1.128,\"\")", PCT),
               ("LCS = LC + 3σ", "=IFERROR(G7+3*G9,\"\")", PCT),
               ("LCI = LC − 3σ", "=IFERROR(G7-3*G9,\"\")", PCT),
               ("Meta ponderada", "=IFERROR(SUMIFS({mk},{est},\"Válida\",{g},{cg},{p},{cp})/SUMIFS({kc},{est},\"Válida\",{g},{cg},{p},{cp}),\"\")", PCT)]
        f0 = 16
        a, b = f0, f0 + len(periodos) - 1
        for k, (lbl, fml, fmt) in enumerate(par):
            r = 7 + k
            ws.cell(r, 6, lbl).font = f_bold
            c = ws.cell(r, 7, fml.format(a=a, b=b, mk=R["metakc"], est=R["estado"], g=R["grupo"], cg=crit_g,
                                         p=R["prod"], cp=crit_p, kc=R["kc"]))
            c.number_format, c.font = fmt, f_bold
        ws.cell(7, 8, "d2 = 1,128 es la constante de la carta I-MR para rangos móviles de 2 puntos.").font = f_nota

        cab = ["Período", "Año ISO" if tipo == "sem" else "Año", "Semana" if tipo == "sem" else "Mes", "Órdenes válidas", "Kg producidos", "Kg consumidos", "Rend. met.",
               "Meta", "vs meta (pp)", "MR", "LC", "LCS", "LCI", "Estado", "Órdenes a corregir"]
        encabezado(ws, f0 - 1, cab)
        for i, per in enumerate(periodos):
            r = f0 + i
            if tipo == "mes":
                anio, mes = per
                ws.cell(r, 1, f"{MESES[mes - 1]} {anio}")
                ws.cell(r, 2, anio)
                ws.cell(r, 3, mes)
                cond = f'{R["anio"]},$B{r},{R["mes"]},$C{r}'
            elif tipo == "sem":
                anio, sem = per
                ws.cell(r, 1, etiqueta_semana(anio, sem))
                ws.cell(r, 2, anio)
                ws.cell(r, 3, sem)
                cond = f'{R["anioiso"]},$B{r},{R["sem"]},$C{r}'
            else:
                ws.cell(r, 1, per).number_format = "dd-mm-yyyy"
                ws.cell(r, 2, f"=YEAR(A{r})")
                ws.cell(r, 3, f"=MONTH(A{r})")
                cond = f'{R["fecha"]},$A{r}'
            filt = f'{R["grupo"]},{crit_g},{R["prod"]},{crit_p}'
            ws.cell(r, 4, f'=COUNTIFS({cond},{R["estado"]},"Válida",{filt})')
            ws.cell(r, 5, f'=SUMIFS({R["kp"]},{cond},{R["estado"]},"Válida",{filt})')
            ws.cell(r, 6, f'=SUMIFS({R["kc"]},{cond},{R["estado"]},"Válida",{filt})')
            ws.cell(r, 7, f'=IF(F{r}=0,NA(),E{r}/F{r})')
            ws.cell(r, 8, f'=IFERROR(SUMIFS({R["metakc"]},{cond},{R["estado"]},"Válida",{filt})/F{r},NA())')
            ws.cell(r, 9, f'=IFERROR((G{r}-H{r})*100,"")')
            ws.cell(r, 10, "" if i == 0 else f'=IFERROR(ABS(G{r}-G{r - 1}),"")')
            ws.cell(r, 11, "=$G$7")
            ws.cell(r, 12, "=$G$10")
            ws.cell(r, 13, "=$G$11")
            ws.cell(r, 14, f'=IF(ISNA(G{r}),"Sin producción",IF(L{r}="","Sin límites",IF(G{r}>L{r},"▲ Sobre LCS",IF(G{r}<M{r},"▼ Bajo LCI","En control"))))')
            ws.cell(r, 15, f'=COUNTIFS({cond},{R["estado"]},"A corregir",{filt})')
            for j in range(1, 16):
                ws.cell(r, j).font = f_base
                ws.cell(r, j).border = borde
            for j in (5, 6):
                ws.cell(r, j).number_format = KG
            for j in (7, 8, 10, 11, 12, 13):
                ws.cell(r, j).number_format = PCT
            ws.cell(r, 9).number_format = PP
        ws.conditional_formatting.add(f"N{a}:N{b}", FormulaRule(formula=[f'LEFT(N{a},1)<>"E"'], fill=fill_alerta,
                                                                font=Font(name=F, bold=True, color="D03B3B")))
        ws.conditional_formatting.add(f"I{a}:I{b}", CellIsRule(operator="lessThan", formula=["0"],
                                                               font=Font(name=F, color="D03B3B")))
        ws.freeze_panes = ws.cell(f0, 2)
        anchos(ws, [20 if tipo == "sem" else 13, 7, 7, 10, 14, 14, 11, 10, 11, 9, 9, 9, 9, 14, 11])
        ws.column_dimensions["F"].width = 26

        ch = LineChart()
        ch.title = titulo_txt
        # ancho proporcional a la cantidad de puntos para que se lea el % de cada uno
        ch.height, ch.width = 10, max(26, len(periodos) * (1.1 if tipo == "dia" else 1.4))
        ch.y_axis.title = "Rend. met."
        ch.y_axis.number_format = "0.0%"
        ch.y_axis.majorGridlines = None
        for col, color, ancho_l, dash in ((7, GRAFITO, 22000, None), (11, "898781", 15000, None),
                                          (12, ROJO, 15000, None), (13, ROJO, 15000, None),
                                          (8, "1BAF7A", 15000, "dash")):
            ch.add_data(Reference(ws, min_col=col, min_row=f0 - 1, max_row=b), titles_from_data=True)
            s = ch.series[-1]
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = ancho_l
            if dash:
                s.graphicalProperties.line.dashStyle = dash
            s.smooth = False
        ch.series[0].marker.symbol = "circle"
        ch.series[0].marker.size = 5
        ch.series[0].dLbls = DataLabelList()
        ch.series[0].dLbls.showVal = True
        ch.series[0].dLbls.showSerName = ch.series[0].dLbls.showCatName = ch.series[0].dLbls.showLegendKey = False
        ch.series[0].dLbls.numFmt = "0.0%"
        ch.series[0].dLbls.position = "t"
        ch.set_categories(Reference(ws, min_col=1, min_row=f0, max_row=b))
        ws.add_chart(ch, "Q4")
        return ws

    meses = sorted({(o["ini"].year, o["ini"].month) for o in ordenes})
    dias = sorted({o["ini"].date() for o in ordenes})
    hoja_carta("Carta mensual", "Carta de control mensual", meses, "mes")
    semanas = sorted({tuple(semana_iso(o["ini"])[:2]) for o in ordenes})
    hoja_carta("Carta semanal", "Carta de control semanal", semanas, "sem")
    hoja_carta("Carta diaria", "Carta de control diaria", dias, "dia")

    # ------------------------------------------------------ Grupo x mes
    wg = wb.create_sheet("Grupo x mes")
    titulo(wg, "Rendimiento metálico por grupo y mes", "Órdenes válidas. Rend. = Σ kg producidos / Σ kg consumidos.")
    anios = sorted({a for a, _ in meses})

    def bloque(fila0, nombre, fn_celda, fmt, total_fn):
        wg.cell(fila0, 1, nombre).font = f_bold
        encabezado(wg, fila0 + 1, ["Grupo"] + [f"{MESES[m - 1]} {a}" for a, m in meses] + ["Total"])
        for i, g in enumerate(GRUPOS + ["Total"]):
            r = fila0 + 2 + i
            wg.cell(r, 1, g).font = f_bold if g == "Total" else f_base
            crit = '"*"' if g == "Total" else f"$A{r}"
            for j, (a, m) in enumerate(meses):
                c = wg.cell(r, 2 + j, fn_celda(crit, f'{R["anio"]},{a},{R["mes"]},{m}'))
                c.number_format, c.font = fmt, f_bold if g == "Total" else f_base
            c = wg.cell(r, 2 + len(meses), total_fn(crit))
            c.number_format, c.font = fmt, f_bold
        return fila0 + 2, fila0 + 2 + len(GRUPOS)

    def rend(crit, cond):
        return (f'=IFERROR(SUMIFS({R["kp"]},{cond},{R["grupo"]},{crit},{R["estado"]},"Válida")/'
                f'SUMIFS({R["kc"]},{cond},{R["grupo"]},{crit},{R["estado"]},"Válida"),"")')

    def rend_tot(crit):
        return rend(crit, f'{R["anio"]},">={anios[0]}"')

    def kgp(crit, cond):
        return f'=SUMIFS({R["kp"]},{cond},{R["grupo"]},{crit},{R["estado"]},"Válida")'

    r0, r1 = bloque(4, "Rendimiento metálico", rend, PCT, rend_tot)
    ultima_col = get_column_letter(1 + len(meses))
    wg.conditional_formatting.add(f"B{r0}:{ultima_col}{r1 - 1}", ColorScaleRule(
        start_type="num", start_value=0.90, start_color="E34948", mid_type="num", mid_value=0.945,
        mid_color="F0EFEC", end_type="num", end_value=0.97, end_color="2A78D6"))
    k0 = r1 + 3
    bloque(k0, "Kg producidos", kgp, KG, lambda crit: kgp(crit, f'{R["anio"]},">={anios[0]}"'))
    m0 = k0 + len(GRUPOS) + 5

    def meta(crit, cond):
        return (f'=IFERROR(SUMIFS({R["metakc"]},{cond},{R["grupo"]},{crit},{R["estado"]},"Válida")/'
                f'SUMIFS({R["kc"]},{cond},{R["grupo"]},{crit},{R["estado"]},"Válida"),"")')
    bloque(m0, "Meta ponderada (Rend_Met × kg consumidos)", meta, PCT,
           lambda crit: meta(crit, f'{R["anio"]},">={anios[0]}"'))
    wg.cell(r1 + 1, 1, "Escala de color: rojo ≤ 90 %, gris 94,5 %, azul ≥ 97 %.").font = f_nota
    anchos(wg, [16] + [11] * (len(meses) + 1))
    wg.freeze_panes = "B6"

    # ------------------------------------------------------ Grupo x semana
    wsg = wb.create_sheet("Grupo x semana")
    titulo(wsg, "Rendimiento metálico por grupo y semana ISO",
           "Semana ISO: lunes a domingo, según la fecha de inicio de la orden. Órdenes válidas. "
           "Rend. = Σ kg producidos / Σ kg consumidos.")
    gcols = GRUPOS + ["Total"]
    encabezado(wsg, 4, ["Semana", "Año ISO", "N° semana"] + [f"Rend. {g}" for g in gcols]
               + ["Meta total", "vs meta (pp)", "Kg producidos", "Kg consumidos", "Órdenes válidas"])
    nc = 3 + len(gcols)
    for i, (y, w) in enumerate(semanas):
        r = 5 + i
        wsg.cell(r, 1, etiqueta_semana(y, w))
        wsg.cell(r, 2, y)
        wsg.cell(r, 3, w)
        cond = f'{R["anioiso"]},$B{r},{R["sem"]},$C{r},{R["estado"]},"Válida"'
        for j, g in enumerate(gcols):
            crit = '"*"' if g == "Total" else f'"{g}"'
            c = wsg.cell(r, 4 + j, f'=IFERROR(SUMIFS({R["kp"]},{cond},{R["grupo"]},{crit})/'
                                    f'SUMIFS({R["kc"]},{cond},{R["grupo"]},{crit}),"")')
            c.number_format = PCT
        tot = get_column_letter(nc)
        wsg.cell(r, nc + 1, f'=IFERROR(SUMIFS({R["metakc"]},{cond})/SUMIFS({R["kc"]},{cond}),"")').number_format = PCT
        wsg.cell(r, nc + 2, f'=IFERROR(({tot}{r}-{get_column_letter(nc + 1)}{r})*100,"")').number_format = PP
        wsg.cell(r, nc + 3, f'=SUMIFS({R["kp"]},{cond})').number_format = KG
        wsg.cell(r, nc + 4, f'=SUMIFS({R["kc"]},{cond})').number_format = KG
        wsg.cell(r, nc + 5, f'=COUNTIFS({cond})')
        for j in range(1, nc + 6):
            wsg.cell(r, j).font = f_bold if j == nc else f_base
            wsg.cell(r, j).border = borde
    us = 4 + len(semanas)
    wsg.conditional_formatting.add(f"D5:{get_column_letter(nc)}{us}", ColorScaleRule(
        start_type="num", start_value=0.90, start_color="E34948", mid_type="num", mid_value=0.945,
        mid_color="F0EFEC", end_type="num", end_value=0.97, end_color="2A78D6"))
    wsg.conditional_formatting.add(f"{get_column_letter(nc + 2)}5:{get_column_letter(nc + 2)}{us}",
                                   CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="D03B3B")))
    wsg.cell(us + 2, 1, "Escala de color: rojo ≤ 90 %, gris 94,5 %, azul ≥ 97 %. Celda vacía = el grupo no produjo esa semana.").font = f_nota
    anchos(wsg, [20, 8, 9] + [11] * len(gcols) + [10, 11, 14, 14, 10])
    wsg.freeze_panes = "D5"

    # ------------------------------------------------------ Ranking productos
    wr = wb.create_sheet("Ranking productos")
    titulo(wr, "Ranking de productos (de mayor a menor kg producidos)",
           "Órdenes válidas. El orden es el del momento de la exportación; usa los filtros del encabezado para reordenar.")
    cab = ["#", "Grupo", "Producto", "Medidas", "Formato (maestro)", "Largo (m)", "Kg producidos", "Kg consumidos",
           "Rend. met.", "Meta", "vs meta (pp)", "Órdenes válidas", "Órdenes a corregir", "Pérdida (kg)"]
    encabezado(wr, 4, cab)
    kg_val = {}
    for o in ordenes:
        if o["kp"] <= o["kc"]:
            kg_val[o["p"]] = kg_val.get(o["p"], 0) + o["kp"]
    orden_prod = sorted(productos, key=lambda p: -kg_val.get(p, 0))
    for i, pn in enumerate(orden_prod):
        r = 5 + i
        p = productos[pn]
        c = f'{R["prod"]},$C{r},{R["estado"]},"Válida"'
        vals = [i + 1, p["grupo"], pn, p["desc"], p.get("formato"), p["largo"],
                f'=SUMIFS({R["kp"]},{c})', f'=SUMIFS({R["kc"]},{c})', f'=IFERROR(G{r}/H{r},"")',
                f'=IFERROR(SUMIFS({R["metakc"]},{c})/H{r},"")', f'=IFERROR((I{r}-J{r})*100,"")',
                f'=COUNTIFS({c})', f'=COUNTIFS({R["prod"]},$C{r},{R["estado"]},"A corregir")', f"=H{r}-G{r}"]
        for j, v in enumerate(vals, 1):
            cc = wr.cell(r, j, v)
            cc.font = f_base
        for j in (7, 8, 14):
            wr.cell(r, j).number_format = KG
        for j in (9, 10):
            wr.cell(r, j).number_format = PCT
        wr.cell(r, 11).number_format = PP
    ur = 4 + len(orden_prod)
    t2 = Table(displayName="Ranking", ref=f"A4:N{ur}")
    t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    wr.add_table(t2)
    wr.conditional_formatting.add(f"K5:K{ur}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, color="D03B3B")))
    wr.conditional_formatting.add(f"M5:M{ur}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill_alerta))
    anchos(wr, [5, 13, 44, 22, 16, 9, 14, 14, 11, 10, 11, 10, 10, 12])
    wr.freeze_panes = "D5"

    # ------------------------------------------------------ A corregir
    wc = wb.create_sheet("A corregir")
    titulo(wc, "Órdenes a corregir (rendimiento > 100 %)",
           "Excluidas de todos los cálculos: producir más kg de los consumidos no es físicamente posible "
           "(consumo mal imputado o ajuste de mes).")
    cab = ["Orden", "Fecha inicio", "Grupo", "Producto", "Kg producidos", "Kg consumidos", "Diferencia (kg)",
           "Rend. met.", "Comentario / corrección"]
    encabezado(wc, 4, cab)
    exc = sorted([o for o in ordenes if o["kp"] > o["kc"]], key=lambda o: -o["kp"] / o["kc"])
    for i, o in enumerate(exc):
        r = 5 + i
        vals = [int(o["o"]) if o["o"].isdigit() else o["o"], o["ini"], productos[o["p"]]["grupo"], o["p"],
                o["kp"], o["kc"], f"=E{r}-F{r}", f"=E{r}/F{r}", None]
        for j, v in enumerate(vals, 1):
            wc.cell(r, j, v).font = f_base
        wc.cell(r, 2).number_format = "dd-mm-yyyy hh:mm"
        for j in (5, 6, 7):
            wc.cell(r, j).number_format = KG
        wc.cell(r, 8).number_format = PCT
        wc.cell(r, 8).font = Font(name=F, bold=True, color="D03B3B")
        wc.cell(r, 9).fill = fill_input
    wc.cell(5 + len(exc) + 1, 1, "Columna amarilla: para anotar la corrección de cada orden.").font = f_nota
    anchos(wc, [11, 16, 13, 44, 13, 13, 14, 11, 40])
    wc.freeze_panes = "A5"

    # ------------------------------------------------------ Productos
    wp = wb.create_sheet("Productos")
    titulo(wp, "Catálogo de productos del período", "Grupo y largo desde el maestro; medidas desde el nombre (mm).")
    cab = ["Producto", "Grupo", "Medidas", "Diámetro (mm)", "Ancho (mm)", "Espesor (mm)", "Lado (mm)",
           "Pulgadas", "Largo (m)", "Largo asumido", "Formato (maestro)", "Código (maestro)", "En maestro"]
    encabezado(wp, 4, cab)
    for i, pn in enumerate(sorted(productos, key=lambda x: (GRUPOS.index(productos[x]["grupo"])
                                                           if productos[x]["grupo"] in GRUPOS else 99, x))):
        p = productos[pn]
        vals = [pn, p["grupo"], p["desc"], p["diametro"], p["ancho"], p["espesor"], p["lado"], p["pulgadas"],
                p["largo"], "Sí" if p["largo_asumido"] else "No", p.get("formato"), p.get("codigo"),
                "No" if pn in sin_maestro else "Sí"]
        for j, v in enumerate(vals, 1):
            wp.cell(5 + i, j, v).font = f_base
    up = 4 + len(productos)
    t3 = Table(displayName="Productos", ref=f"A4:M{up}")
    t3.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    wp.add_table(t3)
    wp.conditional_formatting.add(f"M5:M{up}", CellIsRule(operator="equal", formula=['"No"'], fill=fill_alerta))
    anchos(wp, [44, 13, 24, 11, 10, 10, 9, 10, 9, 10, 16, 14, 10])
    wp.freeze_panes = "B5"

    # ------------------------------------------------------ Resumen / Léeme
    wl = wb.create_sheet("Resumen", 0)
    titulo(wl, "Rendimiento metálico", f"Fuente: {ruta.name}" + (f" + {ruta_maestro.name}" if maestro else "")
           + f" · generado {datetime.now():%d-%m-%Y %H:%M}")
    kpis = [
        ("Rendimiento metálico (ponderado)", f'=SUMIFS({R["kp"]},{R["estado"]},"Válida")/SUMIFS({R["kc"]},{R["estado"]},"Válida")', PCT),
        ("Meta ponderada", f'=SUMIFS({R["metakc"]},{R["estado"]},"Válida")/SUMIFS({R["kc"]},{R["estado"]},"Válida")', PCT),
        ("Diferencia vs meta (pp)", "=(B4-B5)*100", PP),
        ("Kg producidos", f'=SUMIFS({R["kp"]},{R["estado"]},"Válida")', KG),
        ("Kg consumidos", f'=SUMIFS({R["kc"]},{R["estado"]},"Válida")', KG),
        ("Pérdida metálica (kg)", "=B8-B7", KG),
        ("Órdenes válidas", f'=COUNTIFS({R["estado"]},"Válida")', "0"),
        ("Órdenes a corregir (> 100 %)", f'=COUNTIFS({R["estado"]},"A corregir")', "0"),
    ]
    for i, (lbl, fml, fmt) in enumerate(kpis):
        wl.cell(4 + i, 1, lbl).font = f_bold
        c = wl.cell(4 + i, 2, fml)
        c.number_format, c.font = fmt, Font(name=F, size=12, bold=True)
    notas = [
        "Hojas",
        "• Ordenes: todas las órdenes del CSV, con grupo, medidas (mm), largo (m), rendimiento y estado. Base de todo el libro.",
        "• Carta mensual / Carta semanal / Carta diaria: carta de control I-MR con gráfico y el % de cada punto. Las celdas amarillas filtran por grupo o producto.",
        "• Grupo x semana: rendimiento de cada grupo por semana ISO (lunes a domingo), con meta y kg.",
        "• Grupo x mes: rendimiento, kg producidos y meta por grupo y mes.",
        "• Ranking productos: productos de mayor a menor kg producidos.",
        "• A corregir: órdenes sobre 100 %, excluidas de los cálculos.",
        "• Productos: catálogo con medidas y datos del maestro.",
        "",
        "Reglas",
        "• Rend. met. = kg producidos / kg consumidos (sin despunte ni laminilla).",
        "• Rendimiento de un período o grupo = Σ kg producidos / Σ kg consumidos de las órdenes válidas.",
        "• Órdenes > 100 % se marcan «A corregir» y quedan fuera. Si se corrigen los kg en «Ordenes», todo se recalcula.",
        "• Meta = columna Rend_Met del CSV, ponderada por kg consumidos.",
        "• Carta I-MR: LC = rendimiento ponderado, σ = MR̄ / 1,128, LCS/LCI = LC ± 3σ.",
        "• Grupo y largo desde el maestro (Tipo, CORTE T4); si no está, se deducen del nombre. Barra traspaso = Redondo liso; SEMI = mismo producto.",
        "• Pulgadas convertidas a mm (× 25,4). Fecha de la orden = Fecha_inicio.",
        "",
        "Leyenda de colores: celda amarilla con texto azul = dato editable; fondo rosado = requiere atención.",
    ]
    for i, t in enumerate(notas):
        c = wl.cell(14 + i, 1, t)
        c.font = f_bold if t in ("Hojas", "Reglas") else f_base
    if sin_maestro:
        r = 14 + len(notas) + 1
        wl.cell(r, 1, "Productos no encontrados en el maestro (clasificados por nombre):").font = f_bold
        for i, pn in enumerate(sorted(sin_maestro)):
            wl.cell(r + 1 + i, 1, f"• {pn} → {productos[pn]['grupo']}").font = f_base
    anchos(wl, [34, 18])

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title in ("Ordenes",)
        fuente_arial(ws)
    wb["Resumen"].sheet_properties.tabColor = ROJO

    wb.calculation.fullCalcOnLoad = True  # Excel recalcula todo al abrir
    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"{len(ordenes)} órdenes, {len(productos)} productos, {len(exc)} a corregir -> {SALIDA}")


if __name__ == "__main__":
    main()
